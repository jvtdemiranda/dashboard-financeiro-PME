"""
Gera o dashboard financeiro em Excel a partir de data/processed/transacoes_tratadas.csv.

Terceira etapa do pipeline:
1. gerar_dados.py     -> gera data/raw/transacoes_brutas.csv (dado sujo)
2. tratar_dados.py    -> gera data/processed/transacoes_tratadas.csv (dado limpo)
3. gerar_dashboard.py -> gera dashboard/dashboard_financeiro.xlsx (este script)

Os KPIs e as tabelas de resumo (Fluxo Mensal, Contas a Pagar/Receber, DRE)
são escritos como fórmulas do Excel (SUMIFS) que apontam para a aba "Dados" —
se alguém editar um valor de uma transação diretamente na planilha, os totais
recalculam sozinhos. As listas de detalhamento de "Contas a Pagar/Receber"
são um extrato estático tirado do CSV tratado no momento em que este script
roda; para atualizá-las depois de uma mudança na base, é só rodar o script
de novo (mesma lógica de reprodutibilidade dos outros dois scripts).
"""

import io
import os
import re
import shutil
import zipfile
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.chart.data_source import AxDataSource, NumData, NumRef, NumVal, StrData, StrRef, StrVal
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_to_tuple

# --- Paleta: a mesma do painel publicado (dashboard_template.html) ---
BLUE = "2A78D6"        # Entradas / receita
RED = "E34948"         # Saídas / despesa
WARNING = "C87E00"     # status: pendente (texto; o amarelo claro some no fundo branco)
WARNING_BG = "FEF3DD"
CRITICAL = "D03B3B"    # status: atrasado
CRITICAL_BG = "FCEBEA"
SUCCESS_TEXT = "006300"  # status: pago / saldo positivo
SUCCESS_BG = "E6F4E6"
INK = "0B0B0B"
INK_SEC = "52514E"
INK_MUTED = "898781"
GRID = "E1E0D9"
SURFACE = "FCFCFB"
SURFACE_2 = "F2F1EC"
WHITE = "FFFFFF"

# Uma fonte só, em toda célula. Célula com Font() sem nome fica com a
# fonte padrão de cada programa (no LibreOffice, uma serifada) — era isso
# que misturava fontes na planilha.
FONTE = "Arial"

CUR_FMT = '"R$" #,##0.00'
DATE_FMT = "dd/mm/yyyy"
MONTH_FMT = "mmm/yyyy"
PCT_FMT = "0.0%"
MESES_PT = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]

CATEGORIAS_DESPESA = [
    "Aluguel", "Fornecedor", "Imposto", "Manutenção", "Marketing", "Salário",
]

THIN = Side(style="thin", color=GRID)

# Aba "Dados": coluna A é margem; título nas linhas 1-2, cabeçalho na 4,
# transações a partir da 5. Todas as fórmulas das outras abas apontam
# para estas colunas.
COL_DADOS = {"data": "B", "descricao": "C", "categoria": "D", "tipo": "E", "valor": "F", "status": "G"}
CAB_DADOS = 4
PRIMEIRA_DADOS = CAB_DADOS + 1


def faixa_dados(campo, ultima):
    letra = COL_DADOS[campo]
    return f"Dados!${letra}${PRIMEIRA_DADOS}:${letra}${ultima}"


def mes_pt(data):
    return f"{MESES_PT[data.month - 1]}/{data.year}"


def caminhos():
    pasta_raiz = os.path.join(os.path.dirname(__file__), "..")
    csv_tratado = os.path.join(pasta_raiz, "data", "processed", "transacoes_tratadas.csv")
    pasta_dashboard = os.path.join(pasta_raiz, "dashboard")
    os.makedirs(pasta_dashboard, exist_ok=True)
    xlsx_saida = os.path.join(pasta_dashboard, "dashboard_financeiro.xlsx")
    return csv_tratado, xlsx_saida


# Resultado de cada fórmula, calculado aqui em Python com a mesma regra da
# fórmula: {(aba, "B5"): 1234.5}. O openpyxl grava a fórmula sem o resultado,
# e visualizadores que não calculam (celular, prévias online) mostrariam as
# células vazias. salvar_deterministico() grava esses resultados no arquivo.
RESULTADOS = {}


def formula(ws, row, col, expr, resultado):
    c = ws.cell(row=row, column=col, value=expr)
    RESULTADOS[(ws.title, c.coordinate)] = float(resultado)
    return c


def _valores(wb, ref):
    """Valores das células de uma referência de gráfico, usando o resultado das fórmulas."""
    if ":" not in ref.split("!")[-1]:
        ref = f"{ref}:{ref.split('!')[-1]}"
    aba, (c1, l1, c2, l2) = range_to_tuple(ref)
    ws = wb[aba]
    saida = []
    for linha in range(l1, l2 + 1):
        for coluna in range(c1, c2 + 1):
            celula = ws.cell(row=linha, column=coluna)
            saida.append(RESULTADOS.get((aba, celula.coordinate), celula.value))
    return saida, ws.cell(row=l1, column=c1).number_format


def preencher_caches(wb):
    """
    Grava dentro de cada gráfico uma cópia dos dados que ele desenha (é o que
    o Excel faz ao salvar). O openpyxl grava só a referência às células, e
    visualizadores que não calculam mostrariam o gráfico em branco.
    """
    for ws in wb.worksheets:
        for grafico in ws._charts:
            for serie in grafico.series:
                valores, fmt = _valores(wb, serie.val.numRef.f)
                serie.val.numRef.numCache = NumData(
                    formatCode=fmt, ptCount=len(valores),
                    pt=[NumVal(idx=i, v=float(v)) for i, v in enumerate(valores) if v is not None])
                if serie.tx is not None and serie.tx.strRef is not None:
                    nome, _ = _valores(wb, serie.tx.strRef.f)
                    serie.tx.strRef.strCache = StrData(ptCount=1, pt=[StrVal(idx=0, v=str(nome[0]))])
                if serie.cat is not None:
                    ref = serie.cat.numRef.f if serie.cat.numRef is not None else serie.cat.strRef.f
                    rotulos, fmt = _valores(wb, ref)
                    if all(isinstance(r, datetime) for r in rotulos):
                        serie.cat = AxDataSource(numRef=NumRef(f=ref, numCache=NumData(
                            formatCode=fmt, ptCount=len(rotulos),
                            pt=[NumVal(idx=i, v=float((r - datetime(1899, 12, 30)).days)) for i, r in enumerate(rotulos)])))
                    else:
                        serie.cat = AxDataSource(strRef=StrRef(f=ref, strCache=StrData(
                            ptCount=len(rotulos), pt=[StrVal(idx=i, v=str(r).strip()) for i, r in enumerate(rotulos)])))


def eixos_visiveis(chart):
    """
    O openpyxl 3.1 não grava que os eixos são visíveis, e há relatos de
    versões recentes do Excel escondendo o eixo nesse caso (gráfico sem
    meses e sem valores). Marcar explicitamente não muda nada onde já
    funcionava.
    """
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    return chart


def estilo(c, tam=10, negrito=False, italico=False, cor=INK, fmt=None, alinhar=None, fundo=None,
           linha=True, recuo=0, quebra=False):
    """Aplica o estilo da planilha a uma célula: fonte única, fio fino embaixo, recuo."""
    c.font = Font(name=FONTE, size=tam, bold=negrito, italic=italico, color=cor)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=alinhar, vertical="center", indent=recuo, wrap_text=quebra)
    if fundo:
        c.fill = PatternFill(start_color=fundo, end_color=fundo, fill_type="solid")
    c.border = Border(bottom=THIN) if linha else Border()
    return c


def texto(ws, r, col, valor, **kw):
    """Texto vindo dos dados é sempre texto: nunca vira fórmula, mesmo começando com "="."""
    c = ws.cell(row=r, column=col, value=None if valor is None else str(valor))
    c.data_type = "s"
    return estilo(c, **kw)


def larguras(ws, valores):
    """Coluna A é sempre uma margem estreita; `valores` começam na coluna B."""
    ws.column_dimensions["A"].width = 2
    for j, w in enumerate(valores, start=2):
        ws.column_dimensions[get_column_letter(j)].width = w


def titulo_pagina(ws, texto_titulo, subtitulo):
    ws.cell(row=1, column=2, value=texto_titulo).font = Font(name=FONTE, size=17, bold=True, color=INK)
    ws.cell(row=2, column=2, value=subtitulo).font = Font(name=FONTE, size=10, italic=True, color=INK_SEC)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 10
    ws.sheet_view.showGridLines = False


def secao(ws, r, nome, col1, col2):
    """Título de seção com um traço embaixo, cobrindo as colunas da tabela."""
    for col in range(col1, col2 + 1):
        ws.cell(row=r, column=col).border = Border(bottom=Side(style="medium", color=INK))
    ws.cell(row=r, column=col1, value=nome).font = Font(name=FONTE, size=12, bold=True, color=INK)
    ws.row_dimensions[r].height = 22


def header_row(ws, row, nomes, col_inicial=2, numericas=(), centro=()):
    """Cabeçalho escuro; colunas de valor alinhadas à direita, como os números."""
    fill = PatternFill(start_color=INK, end_color=INK, fill_type="solid")
    for j, nome in enumerate(nomes):
        cell = ws.cell(row=row, column=col_inicial + j, value=nome)
        cell.fill = fill
        cell.font = Font(name=FONTE, size=10, bold=True, color=WHITE)
        alinhar = "right" if j in numericas else "center" if j in centro else "left"
        cell.alignment = Alignment(horizontal=alinhar, vertical="center", indent=1 if alinhar == "left" else 0)
    ws.row_dimensions[row].height = 26


def cores_de_status(ws, faixa):
    """
    Status como "etiqueta" colorida, igual ao painel: Pago verde, Pendente
    âmbar, Atrasado vermelho. Por formatação condicional, para a cor seguir
    o valor se alguém editar o status na planilha.
    """
    for valor, cor, fundo in [("Pago", SUCCESS_TEXT, SUCCESS_BG), ("Pendente", WARNING, WARNING_BG),
                              ("Atrasado", CRITICAL, CRITICAL_BG)]:
        ws.conditional_formatting.add(faixa, CellIsRule(
            operator="equal", formula=[f'"{valor}"'],
            font=Font(name=FONTE, bold=True, color=cor),
            fill=PatternFill(start_color=fundo, end_color=fundo, fill_type="solid")))


def cores_de_saldo(ws, faixa, negrito=False, tam=10):
    ws.conditional_formatting.add(faixa, CellIsRule(
        operator="lessThan", formula=["0"], font=Font(name=FONTE, bold=negrito, size=tam, color=CRITICAL)))
    ws.conditional_formatting.add(faixa, CellIsRule(
        operator="greaterThanOrEqual", formula=["0"], font=Font(name=FONTE, bold=negrito, size=tam, color=SUCCESS_TEXT)))


def montar_dados(wb, df):
    ws = wb.create_sheet("Dados")
    titulo_pagina(ws, "Transações",
                  f"{len(df)} lançamentos já tratados — a base de todas as fórmulas desta planilha: "
                  "mude um valor ou um status aqui e os totais das outras abas recalculam.")
    colunas = ["Data", "Descrição", "Categoria", "Tipo", "Valor", "Status"]
    header_row(ws, CAB_DADOS, colunas, numericas=(4,), centro=(5,))

    for i, row in enumerate(df.itertuples(index=False), start=PRIMEIRA_DADOS):
        estilo(ws.cell(row=i, column=2, value=row.data.to_pydatetime()), fmt=DATE_FMT, cor=INK_SEC, recuo=1,
               alinhar="left")
        texto(ws, i, 3, row.descricao, negrito=True, recuo=1)
        texto(ws, i, 4, row.categoria, cor=INK_SEC, recuo=1)
        texto(ws, i, 5, row.tipo, cor=BLUE if row.tipo == "Entrada" else RED, recuo=1)
        estilo(ws.cell(row=i, column=6, value=float(row.valor)), fmt=CUR_FMT)
        texto(ws, i, 7, row.status, alinhar="center")
        ws.row_dimensions[i].height = 18

    ultima = PRIMEIRA_DADOS + len(df) - 1
    cores_de_status(ws, f"G{PRIMEIRA_DADOS}:G{ultima}")
    larguras(ws, [13, 30, 20, 11, 15, 13])
    ws.freeze_panes = f"A{PRIMEIRA_DADOS}"
    ws.auto_filter.ref = f"B{CAB_DADOS}:G{ultima}"
    ws.print_title_rows = f"{CAB_DADOS}:{CAB_DADOS}"
    return ws, ultima


def montar_fluxo_mensal(wb, df, meses, u):
    ws = wb.create_sheet("Fluxo Mensal")
    titulo_pagina(ws, "Fluxo de caixa mensal",
                  f"{mes_pt(meses[0])} a {mes_pt(meses[-1])} · regime de caixa: só transações com status Pago "
                  "(pendentes e atrasadas estão em Contas a Pagar e Receber)")

    header_row(ws, 4, ["Mês", "Entradas", "Saídas", "Saldo do mês", "Saldo acumulado"], numericas=(1, 2, 3, 4))

    primeira_linha = 5
    pagos = df[df["status"] == "Pago"]
    acumulado = 0.0
    totais = {"C": 0.0, "D": 0.0, "E": 0.0}
    data, tipo, valor, status = (faixa_dados(c, u) for c in ("data", "tipo", "valor", "status"))
    for i, mes in enumerate(meses):
        r = primeira_linha + i
        estilo(ws.cell(row=r, column=2, value=mes), fmt=MONTH_FMT, negrito=True, recuo=1, alinhar="left")
        no_mes = pagos[(pagos["data"] >= mes) & (pagos["data"] < mes + pd.DateOffset(months=1))]
        valores = {}
        for col, t, cor in [(3, "Entrada", BLUE), (4, "Saída", RED)]:
            valores[col] = no_mes.loc[no_mes["tipo"] == t, "valor"].sum()
            estilo(formula(ws, r, col,
                           f'=SUMIFS({valor},{tipo},"{t}",{status},"Pago",'
                           f'{data},">="&B{r},{data},"<"&EDATE(B{r},1))',
                           valores[col]), fmt=CUR_FMT, cor=cor)
        saldo = valores[3] - valores[4]
        acumulado += saldo
        estilo(formula(ws, r, 5, f"=C{r}-D{r}", saldo), fmt=CUR_FMT, negrito=True)
        estilo(formula(ws, r, 6, f"=E{r}" if i == 0 else f"=F{r - 1}+E{r}", acumulado), fmt=CUR_FMT)
        ws.row_dimensions[r].height = 20
        totais["C"] += valores[3]
        totais["D"] += valores[4]
        totais["E"] += saldo

    ultima_linha = primeira_linha + len(meses) - 1
    linha_total = ultima_linha + 1
    borda_total = Border(top=Side(style="medium", color=INK))
    c = estilo(ws.cell(row=linha_total, column=2, value="Total"), negrito=True, recuo=1, linha=False)
    c.border = borda_total
    for col, letra in [(3, "C"), (4, "D"), (5, "E")]:
        cell = formula(ws, linha_total, col, f"=SUM({letra}{primeira_linha}:{letra}{ultima_linha})", totais[letra])
        estilo(cell, negrito=True, fmt=CUR_FMT, linha=False).border = borda_total
    ws.cell(row=linha_total, column=6).border = borda_total
    ws.row_dimensions[linha_total].height = 22

    cores_de_saldo(ws, f"E{primeira_linha}:E{ultima_linha}", negrito=True)
    cores_de_saldo(ws, f"E{linha_total}", negrito=True)
    larguras(ws, [14, 17, 17, 17, 20])

    # Gráficos logo abaixo da tabela, na largura dela (antes ficavam longe, à direita)
    categorias = Reference(ws, min_col=2, min_row=primeira_linha, max_row=ultima_linha)
    linha_graficos = linha_total + 3
    secao(ws, linha_graficos - 1, "Entradas x saídas por mês", 2, 6)

    chart1 = eixos_visiveis(BarChart())
    chart1.type = "col"
    chart1.grouping = "clustered"
    chart1.title = None
    chart1.style = 10
    chart1.y_axis.title = None
    chart1.y_axis.numFmt = '"R$" #,##0'
    chart1.x_axis.title = None
    chart1.gapWidth = 60
    dados = Reference(ws, min_col=3, max_col=4, min_row=4, max_row=ultima_linha)
    chart1.add_data(dados, titles_from_data=True)
    chart1.set_categories(categorias)
    chart1.series[0].graphicalProperties.solidFill = BLUE
    chart1.series[1].graphicalProperties.solidFill = RED
    chart1.legend.position = "b"
    chart1.height = 8
    chart1.width = 20.5
    ws.add_chart(chart1, f"B{linha_graficos}")

    linha_saldo = linha_graficos + 18
    secao(ws, linha_saldo - 1, "Saldo acumulado", 2, 6)
    chart2 = eixos_visiveis(LineChart())
    chart2.title = None
    chart2.style = 10
    chart2.y_axis.title = None
    chart2.y_axis.numFmt = '"R$" #,##0'
    dados2 = Reference(ws, min_col=6, min_row=4, max_row=ultima_linha)
    chart2.add_data(dados2, titles_from_data=True)
    chart2.set_categories(categorias)
    s = chart2.series[0]
    s.graphicalProperties.line.solidFill = BLUE
    s.graphicalProperties.line.width = 22000
    s.marker = Marker(symbol="circle", size=6)
    s.marker.graphicalProperties.solidFill = BLUE
    s.marker.graphicalProperties.line.solidFill = BLUE
    s.smooth = False
    chart2.legend = None
    chart2.height = 8
    chart2.width = 20.5
    ws.add_chart(chart2, f"B{linha_saldo}")

    return ws, primeira_linha, ultima_linha, linha_total


def montar_contas(wb, df, u):
    ws = wb.create_sheet("Contas a Pagar e Receber")
    titulo_pagina(ws, "Contas a pagar e receber",
                  "Transações com status Pendente ou Atrasado — o que ainda vai entrar e sair do caixa")

    # o resumo usa a mesma grade das listas abaixo: rótulo em B:C, números em D, E e F
    header_row(ws, 4, ["Em aberto", "", "Pendente", "Atrasado", "Total"], numericas=(2, 3, 4))
    ws.merge_cells("B4:C4")
    tipo, valor, status = (faixa_dados(c, u) for c in ("tipo", "valor", "status"))
    em_aberto = {}
    for r, t, rotulo, cor in [(5, "Entrada", "A receber (entradas)", BLUE), (6, "Saída", "A pagar (saídas)", RED)]:
        estilo(ws.cell(row=r, column=2, value=rotulo), negrito=True, recuo=1)
        estilo(ws.cell(row=r, column=3))
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        soma = 0.0
        for c, st, cor_st in [(4, "Pendente", WARNING), (5, "Atrasado", CRITICAL)]:
            v = df.loc[(df["tipo"] == t) & (df["status"] == st), "valor"].sum()
            soma += v
            estilo(formula(ws, r, c, f'=SUMIFS({valor},{tipo},"{t}",{status},"{st}")', v), fmt=CUR_FMT, cor=cor_st)
        em_aberto[r] = soma
        estilo(formula(ws, r, 6, f"=D{r}+E{r}", soma), fmt=CUR_FMT, negrito=True, cor=cor)
        ws.row_dimensions[r].height = 22

    borda_total = Border(top=Side(style="medium", color=INK))
    c = estilo(ws.cell(row=7, column=2, value="Saldo em aberto (receber − pagar)"), negrito=True, recuo=1, linha=False)
    c.border = borda_total
    for col in (3, 4, 5):
        ws.cell(row=7, column=col).border = borda_total
    ws.merge_cells("B7:C7")
    saldo_cell = formula(ws, 7, 6, "=F5-F6", em_aberto[5] - em_aberto[6])
    estilo(saldo_cell, negrito=True, tam=11, fmt=CUR_FMT, linha=False).border = borda_total
    cores_de_saldo(ws, "F7", negrito=True, tam=11)
    ws.row_dimensions[7].height = 24

    receber = df[(df["tipo"] == "Entrada") & (df["status"].isin(["Pendente", "Atrasado"]))].sort_values("data")
    pagar = df[(df["tipo"] == "Saída") & (df["status"].isin(["Pendente", "Atrasado"]))].sort_values("data")

    def escrever_detalhe(bloco_df, titulo, linha_titulo, cor_valor):
        secao(ws, linha_titulo, f"{titulo} ({len(bloco_df)})", 2, 6)
        linha_cabecalho = linha_titulo + 1
        header_row(ws, linha_cabecalho, ["Data", "Descrição", "Categoria", "Valor", "Status"], numericas=(3,), centro=(4,))
        for i, row in enumerate(bloco_df.itertuples(index=False), start=1):
            r = linha_cabecalho + i
            estilo(ws.cell(row=r, column=2, value=row.data.to_pydatetime()), fmt=DATE_FMT, cor=INK_SEC, recuo=1,
                   alinhar="left")
            texto(ws, r, 3, row.descricao, negrito=True, recuo=1)
            texto(ws, r, 4, row.categoria, cor=INK_SEC, recuo=1)
            estilo(ws.cell(row=r, column=5, value=float(row.valor)), fmt=CUR_FMT, cor=cor_valor)
            texto(ws, r, 6, row.status, alinhar="center")
            ws.row_dimensions[r].height = 18
        linha_final = linha_cabecalho + len(bloco_df)
        cores_de_status(ws, f"F{linha_cabecalho + 1}:F{linha_final}")
        return linha_final

    # uma lista embaixo da outra: lado a lado, ficavam espremidas e cortadas
    fim = escrever_detalhe(receber, "Contas a receber em aberto", 10, BLUE)
    escrever_detalhe(pagar, "Contas a pagar em aberto", fim + 3, RED)

    larguras(ws, [14, 30, 20, 16, 16])
    return ws


def montar_dre(wb, df, u):
    ws = wb.create_sheet("DRE")
    titulo_pagina(ws, "DRE simplificado",
                  "Demonstrativo de resultado · regime de competência: todas as transações lançadas, pagas ou não")

    tipo, valor, categoria = (faixa_dados(c, u) for c in ("tipo", "valor", "categoria"))
    receita = df.loc[df["tipo"] == "Entrada", "valor"].sum()
    estilo(ws.cell(row=4, column=2, value="Receita bruta (entradas)"), negrito=True, tam=11, recuo=1)
    estilo(formula(ws, 4, 3, f'=SUMIFS({valor},{tipo},"Entrada")', receita), negrito=True, tam=11, fmt=CUR_FMT, cor=BLUE)
    ws.row_dimensions[4].height = 24

    estilo(ws.cell(row=5, column=2, value="Despesas operacionais"), negrito=True, cor=INK_SEC, recuo=1, linha=False)
    ws.row_dimensions[5].height = 22

    primeira = 6
    total_despesas = 0.0
    for i, cat in enumerate(CATEGORIAS_DESPESA):
        r = primeira + i
        # o rótulo fica sem espaços no início: o gráfico usa esta coluna como nome das barras
        estilo(ws.cell(row=r, column=2, value=cat), cor=INK_SEC, recuo=3)
        v = df.loc[(df["tipo"] == "Saída") & (df["categoria"] == cat), "valor"].sum()
        total_despesas += v
        estilo(formula(ws, r, 3, f'=SUMIFS({valor},{tipo},"Saída",{categoria},"{cat}")', v), fmt=CUR_FMT)
        ws.row_dimensions[r].height = 19
    ultima = primeira + len(CATEGORIAS_DESPESA) - 1

    linha_total_desp = ultima + 1
    estilo(ws.cell(row=linha_total_desp, column=2, value="Total de despesas"), negrito=True, recuo=1)
    estilo(formula(ws, linha_total_desp, 3, f"=SUM(C{primeira}:C{ultima})", total_despesas),
           negrito=True, fmt=CUR_FMT, cor=RED)
    ws.row_dimensions[linha_total_desp].height = 22

    linha_resultado = linha_total_desp + 2
    borda_total = Border(top=Side(style="medium", color=INK))
    c = estilo(ws.cell(row=linha_resultado, column=2, value="Resultado líquido"), negrito=True, tam=12, recuo=1, linha=False)
    c.border = borda_total
    resultado = receita - total_despesas
    estilo(formula(ws, linha_resultado, 3, f"=C4-C{linha_total_desp}", resultado),
           negrito=True, tam=12, fmt=CUR_FMT, linha=False).border = borda_total
    cores_de_saldo(ws, f"C{linha_resultado}", negrito=True, tam=12)
    ws.row_dimensions[linha_resultado].height = 26

    linha_margem = linha_resultado + 1
    estilo(ws.cell(row=linha_margem, column=2, value="Margem líquida"), italico=True, cor=INK_SEC, recuo=1, linha=False)
    estilo(formula(ws, linha_margem, 3, f"=C{linha_resultado}/C4", resultado / receita if receita else 0.0),
           italico=True, cor=INK_SEC, fmt=PCT_FMT, linha=False)

    larguras(ws, [30, 18, 3])

    chart = eixos_visiveis(BarChart())
    chart.type = "bar"
    chart.title = None
    chart.style = 10
    chart.legend = None
    chart.x_axis.title = None
    chart.y_axis.title = None
    chart.y_axis.numFmt = '"R$" #,##0'
    dados = Reference(ws, min_col=3, min_row=primeira, max_row=ultima)
    categorias = Reference(ws, min_col=2, min_row=primeira, max_row=ultima)
    chart.add_data(dados, titles_from_data=False)
    chart.set_categories(categorias)
    chart.series[0].graphicalProperties.solidFill = RED
    chart.height = 8
    chart.width = 16
    secao(ws, 4, "Despesas por categoria", 5, 11)
    ws.add_chart(chart, "E5")

    return ws, linha_resultado


def montar_resumo(wb, meses, linha_total_fluxo, linha_resultado_dre, n_transacoes):
    ws = wb.create_sheet("Resumo", 0)
    titulo_pagina(ws, "Dashboard financeiro — PME",
                  f"{mes_pt(meses[0])} a {mes_pt(meses[-1])} · {n_transacoes} transações tratadas "
                  "a partir de um export bruto simulado (ver README)")

    def tile(r, col, titulo, nota, aba, celula, cor_valor):
        """Cartão de indicador: rótulo, valor e uma nota, numa caixa de 2 colunas."""
        fundo = PatternFill(start_color=SURFACE_2, end_color=SURFACE_2, fill_type="solid")
        for rr in (r, r + 1, r + 2):
            ws.merge_cells(start_row=rr, start_column=col, end_row=rr, end_column=col + 1)
            for cc in (col, col + 1):
                ws.cell(row=rr, column=cc).fill = fundo
        lbl = ws.cell(row=r, column=col, value=titulo)
        lbl.font = Font(name=FONTE, size=9, bold=True, color=INK_SEC)
        lbl.alignment = Alignment(horizontal="left", indent=1, vertical="bottom")
        ref = f"'{aba}'!{celula}" if " " in aba else f"{aba}!{celula}"
        val = formula(ws, r + 1, col, "=" + ref, RESULTADOS[(aba, celula)])
        val.font = Font(name=FONTE, size=17, bold=True, color=cor_valor)
        val.number_format = CUR_FMT
        val.alignment = Alignment(horizontal="left", indent=1, vertical="center")
        n = ws.cell(row=r + 2, column=col, value=nota)
        n.font = Font(name=FONTE, size=9, color=INK_MUTED)
        n.alignment = Alignment(horizontal="left", indent=1, vertical="top")

    for r, h in ((4, 20), (5, 30), (6, 18), (8, 20), (9, 30), (10, 18)):
        ws.row_dimensions[r].height = h
    ws.row_dimensions[7].height = 10

    tile(4, 2, "RECEBIDO", "entrou no caixa (pago)", "Fluxo Mensal", f"C{linha_total_fluxo}", BLUE)
    tile(4, 5, "PAGO", "saiu do caixa (pago)", "Fluxo Mensal", f"D{linha_total_fluxo}", RED)
    tile(4, 8, "SALDO DE CAIXA", "recebido − pago", "Fluxo Mensal", f"E{linha_total_fluxo}", INK)
    tile(8, 2, "A RECEBER", "pendente ou atrasado", "Contas a Pagar e Receber", "F5", BLUE)
    tile(8, 5, "A PAGAR", "pendente ou atrasado", "Contas a Pagar e Receber", "F6", RED)
    tile(8, 8, "RESULTADO LÍQUIDO", "receita − despesas (DRE)", "DRE", f"C{linha_resultado_dre}", INK)

    secao(ws, 12, "Fluxo de caixa mensal", 2, 9)
    chart = eixos_visiveis(BarChart())
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = None
    chart.style = 10
    chart.y_axis.title = None
    chart.y_axis.numFmt = '"R$" #,##0'
    fluxo_ws = wb["Fluxo Mensal"]
    primeira_linha_fluxo = linha_total_fluxo - len(meses)
    ultima_linha_fluxo = linha_total_fluxo - 1
    dados = Reference(fluxo_ws, min_col=3, max_col=4, min_row=4, max_row=ultima_linha_fluxo)
    categorias = Reference(fluxo_ws, min_col=2, min_row=primeira_linha_fluxo, max_row=ultima_linha_fluxo)
    chart.add_data(dados, titles_from_data=True)
    chart.set_categories(categorias)
    chart.series[0].graphicalProperties.solidFill = BLUE
    chart.series[1].graphicalProperties.solidFill = RED
    chart.legend.position = "b"
    chart.height = 8
    chart.width = 22
    ws.add_chart(chart, "B14")

    secao(ws, 31, "Abas desta planilha", 2, 9)
    guia = [
        ("Fluxo Mensal", "entradas e saídas pagas mês a mês, saldo do mês e acumulado"),
        ("Contas a Pagar e Receber", "o que está pendente ou atrasado, com a lista de cada lado"),
        ("DRE", "receita, despesas por categoria e resultado líquido"),
        ("Dados", "as transações tratadas — base de todas as fórmulas; edite e os totais recalculam"),
    ]
    for i, (nome, desc) in enumerate(guia):
        r = 32 + i
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=9)
        estilo(ws.cell(row=r, column=2, value=nome), negrito=True, recuo=1)
        estilo(ws.cell(row=r, column=5, value=desc), cor=INK_SEC, recuo=1)
        for col in range(3, 10):
            if col != 5:
                ws.cell(row=r, column=col).border = Border(bottom=THIN)
        ws.row_dimensions[r].height = 20

    larguras(ws, [15, 15, 2, 15, 15, 2, 15, 15])
    nota = ws.cell(row=37, column=2,
                   value="Dados simulados com sujeira proposital (categorias, datas, sinais e status) — "
                         "o README do projeto mostra o tratamento e o relatório de qualidade de dados.")
    nota.font = Font(name=FONTE, size=8, italic=True, color=INK_MUTED)


def salvar_deterministico(wb, caminho):
    """
    Mesmos dados -> arquivo idêntico, byte a byte. O openpyxl carimba a
    hora atual em cada arquivo interno do .xlsx (um zip) e na data de
    "modificado"; fixar essas datas permite que o CI confira se a cópia
    publicada em public/ está em dia com os dados.
    """
    wb.properties.created = wb.properties.modified = datetime(2026, 1, 1)
    # O Excel recalcula tudo ao abrir; os resultados gravados abaixo servem
    # pra quem abre num visualizador que não calcula fórmulas.
    wb.calculation.fullCalcOnLoad = True
    arquivo_da_aba = {f"xl/worksheets/sheet{i}.xml": ws.title for i, ws in enumerate(wb.worksheets, start=1)}
    gravados = 0
    buffer = io.BytesIO()
    wb.save(buffer)
    with zipfile.ZipFile(buffer) as origem, zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as destino:
        for item in origem.infolist():
            conteudo = origem.read(item.filename)
            aba = arquivo_da_aba.get(item.filename)
            if aba:
                for (nome, celula), valor in RESULTADOS.items():
                    if nome != aba:
                        continue
                    conteudo, n = re.subn(
                        rb'(<c r="' + celula.encode() + rb'"[^>]*>)(<f>.*?</f>)(?:<v\s*/>|<v></v>)',
                        rb"\g<1>\g<2><v>" + repr(round(valor, 10)).encode() + b"</v>", conteudo)
                    gravados += n
            if item.filename == "docProps/core.xml":
                conteudo = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", conteudo)
            info = zipfile.ZipInfo(item.filename, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            destino.writestr(info, conteudo)
    if gravados != len(RESULTADOS):
        raise RuntimeError(f"{len(RESULTADOS) - gravados} fórmula(s) ficaram sem resultado gravado — "
                           "a planilha apareceria vazia em visualizadores que não calculam")


def main():
    csv_tratado, xlsx_saida = caminhos()
    df = pd.read_csv(csv_tratado, parse_dates=["data"])

    meses = pd.date_range(df["data"].min().to_period("M").to_timestamp(),
                           df["data"].max().to_period("M").to_timestamp(), freq="MS")
    meses = list(meses)

    wb = Workbook()
    wb.remove(wb.active)  # remove a aba "Sheet" padrão

    _, ultima_linha_dados = montar_dados(wb, df)
    n_linhas = len(df)
    _, primeira_linha, ultima_linha, linha_total_fluxo = montar_fluxo_mensal(wb, df, meses, ultima_linha_dados)
    montar_contas(wb, df, ultima_linha_dados)
    _, linha_resultado_dre = montar_dre(wb, df, ultima_linha_dados)
    montar_resumo(wb, meses, linha_total_fluxo, linha_resultado_dre, n_linhas)

    # Impressão: cada aba cabe na largura de uma folha A4 deitada (sem isso,
    # tabelas e gráficos saíam cortados entre páginas).
    for ws in wb.worksheets:
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

    wb.active = 0
    preencher_caches(wb)
    salvar_deterministico(wb, xlsx_saida)
    # Cópia em public/ pro botão "Baixar planilha" do painel publicado.
    publico = os.path.join(os.path.dirname(xlsx_saida), "..", "public", "dashboard_financeiro.xlsx")
    shutil.copyfile(xlsx_saida, publico)
    print(f"Dashboard gerado com sucesso: {xlsx_saida}")
    print(f"Linhas de dados: {n_linhas} | Meses no fluxo de caixa: {len(meses)}")


if __name__ == "__main__":
    main()
